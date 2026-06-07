from fastapi import APIRouter, Depends, HTTPException, status, Body, File, Form, UploadFile, Response
from typing import List, Dict, Any, Optional
from sqlalchemy.orm import Session
from app.database import get_db
from app.models import *
from app.schemas import *
from app.main import *
import math
import io
import os

teacher_router = APIRouter(prefix='/api', tags=['teacher'], dependencies=[Depends(RoleChecker(['teacher']))])

@teacher_router.get('/teacher/dashboard')
def teacher_dashboard(current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    return dashboard(current_user, db)

@teacher_router.get('/teacher/courses')
def teacher_courses(current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> list[dict]:
    return build_teacher_course_cards(db, current_user.id)

@teacher_router.post('/teacher/courses')
def teacher_create_course(payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    education_level, grade_level, strand = validate_grade_path(clean_text(payload.get('educationLevel') or payload.get('education_level') or 'Junior High School', max_length=40), clean_text(payload.get('gradeLevel') or payload.get('grade_level') or 'Grade 7', max_length=50), clean_text(payload.get('strand'), required=False, max_length=80))
    code = validate_safe_text(payload.get('code') or f"{(payload.get('subject') or 'CLS')}-{current_user.username[:6]}", 'Class code', max_length=50)
    if db.scalar(select(Course).where(func.lower(Course.code) == code.lower())):
        raise HTTPException(status_code=409, detail='Class code already exists')
    course = Course(code=code, title=validate_safe_text(payload.get('title'), 'Class name', max_length=255), subject=validate_deped_subject(payload.get('subject')), education_level=education_level, grade_level=grade_level, strand=strand, term=validate_term(payload.get('term') or 'Term 1'), school_year=clean_text(payload.get('schoolYear') or payload.get('school_year') or '2026-2027', max_length=20), description=clean_text(payload.get('description') or 'Faculty-created DANILO class.', max_length=1000), teacher_id=current_user.id, is_active=True)
    db.add(course)
    db.flush()
    log_action(db, current_user, 'create_class', 'course', course.id)
    db.commit()
    db.refresh(course)
    return serialize_course(course)

@teacher_router.put('/teacher/courses/{course_id}')
def teacher_update_course(course_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    for attr, key, limit in [('title', 'title', 255), ('school_year', 'schoolYear', 20), ('description', 'description', 1000)]:
        if key in payload:
            setattr(course, attr, validate_safe_text(payload.get(key), key, max_length=limit) if attr == 'title' else clean_text(payload.get(key), max_length=limit))
    if 'subject' in payload:
        course.subject = validate_deped_subject(payload.get('subject'))
    if 'term' in payload:
        course.term = validate_term(payload.get('term'))
    if {'educationLevel', 'gradeLevel', 'strand'} & set(payload.keys()):
        course.education_level, course.grade_level, course.strand = validate_grade_path(clean_text(payload.get('educationLevel') or course.education_level, max_length=40), clean_text(payload.get('gradeLevel') or course.grade_level, max_length=50), clean_text(payload.get('strand') or course.strand, required=False, max_length=80))
    log_action(db, current_user, 'update_class', 'course', course.id)
    db.commit()
    return serialize_course(course)

@teacher_router.get('/teacher/courses/{course_id}/students')
def teacher_course_students(course_id: str, current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> list[dict]:
    ensure_teacher_course(db, current_user, course_id)
    rows = db.execute(select(User, Enrollment).join(Enrollment, Enrollment.student_id == User.id).where(Enrollment.course_id == course_id, Enrollment.status == 'active').order_by(User.full_name)).all()
    return [{**serialize_user(student), 'enrollmentStatus': enrollment.status} for student, enrollment in rows]

@teacher_router.post('/teacher/courses/{course_id}/announcements')
def teacher_create_announcement(course_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    post = StreamPost(course_id=course.id, author_id=current_user.id, title=validate_safe_text(payload.get('title'), 'Title', max_length=255), body=validate_safe_text(payload.get('body'), 'Body', max_length=2000), post_type='announcement')
    db.add(post)
    db.commit()
    return {'ok': True, 'id': post.id}

@teacher_router.get('/teacher/announcements')
def teacher_list_announcements(current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> list[dict]:
    return [item for item in build_stream(db, current_user) if item.get('postType') == 'announcement']

@teacher_router.post('/teacher/announcements')
def teacher_create_announcement_compat(payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course_id = parse_id(payload.get('courseId') or payload.get('course_id'), 'Class')
    return teacher_create_announcement(course_id, payload, current_user, db)

@teacher_router.post('/teacher/courses/{course_id}/modules')
def teacher_create_module(course_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    module = Module(course_id=course.id, melc_code=clean_text(payload.get('melcCode') or payload.get('melc_code') or 'TEACHER-CREATED', max_length=120), learning_competency=clean_text(payload.get('learningCompetency') or payload.get('learning_competency'), required=False, max_length=2000), lesson_objectives=clean_text(payload.get('lessonObjectives') or payload.get('lesson_objectives'), required=False, max_length=2000), assessment_type=validate_assessment_type(payload.get('assessmentType') or payload.get('assessment_type')), grade_level=course.grade_level, subject=course.subject, term=validate_term(payload.get('term') or course.term), week=parse_int(payload.get('week') or 1, 'Week', minimum=1), sequence_order=parse_int(payload.get('sequenceOrder') or payload.get('sequence_order') or 1, 'Sequence order', minimum=1), folder_name=clean_text(payload.get('folderName') or payload.get('folder_name') or f'{course.code}/Lessons', max_length=255), title=clean_text(payload.get('title'), max_length=255), summary=clean_text(payload.get('summary') or 'Teacher-created lesson module.', max_length=2000), essential_question=clean_text(payload.get('essentialQuestion') or payload.get('essential_question') or 'What will you learn from this lesson?', max_length=1000), content=clean_text(payload.get('content'), required=False, max_length=5000), file_url=clean_text(payload.get('fileUrl') or payload.get('file_url'), required=False, max_length=500))
    db.add(module)
    log_action(db, current_user, 'create_module', 'module', None, f'course={course.code}')
    db.commit()
    db.refresh(module)
    try:
        index_module_for_rag(module)
    except Exception:
        logger.exception('Could not index module %s for AI retrieval', module.id)
    return {'ok': True, 'id': module.id}

@teacher_router.post('/teacher/courses/{course_id}/materials/generate')
async def teacher_generate_lesson_from_material(course_id: str, material: UploadFile=File(...), mode: str=Form('normal'), save: bool=True, current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    filename = clean_text(material.filename or 'uploaded-material.txt', max_length=255)
    data = await material.read()
    extracted = extract_material_text(filename, data)
    lesson = await generated_lesson_from_text(course, filename, extracted, mode)
    module_id = None
    if save:
        module = Module(course_id=course.id, melc_code=lesson['melcCode'], learning_competency=lesson['learningCompetency'], lesson_objectives=lesson['lessonObjectives'], assessment_type=lesson['assessmentType'], grade_level=course.grade_level, subject=course.subject, term=lesson['term'], week=lesson['week'], sequence_order=lesson['sequenceOrder'], folder_name=lesson['folderName'], title=lesson['title'], summary=lesson['summary'], essential_question=lesson['essentialQuestion'], content=lesson['content'], file_url=None)
        db.add(module)
        db.add(StreamPost(course_id=course.id, author_id=current_user.id, title=f"New lesson: {lesson['title']}", body='A teacher-editable AI-generated lesson draft was added from uploaded material.', post_type='lesson'))
        log_action(db, current_user, 'generate_lesson_from_material', 'module', None, f'source={filename}')
        db.commit()
        db.refresh(module)
        try:
            index_module_for_rag(module)
        except Exception:
            logger.exception('Could not index generated module %s for AI retrieval', module.id)
        module_id = module.id
    return {'ok': True, 'id': module_id, 'lesson': lesson, 'extractedChars': len(extracted), 'message': 'Lesson generated. Review and edit before teaching.'}

@teacher_router.put('/teacher/modules/{module_id}')
def teacher_update_module(module_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    module = db.get(Module, module_id)
    if not module:
        raise HTTPException(status_code=404, detail='Module not found')
    ensure_teacher_course(db, current_user, module.course_id)
    for attr, key, limit in [('title', 'title', 255), ('summary', 'summary', 2000), ('essential_question', 'essentialQuestion', 1000), ('learning_competency', 'learningCompetency', 2000), ('lesson_objectives', 'lessonObjectives', 2000), ('content', 'content', 5000), ('file_url', 'fileUrl', 500), ('melc_code', 'melcCode', 120), ('folder_name', 'folderName', 255)]:
        if key in payload:
            setattr(module, attr, clean_text(payload.get(key), required=attr not in {'content', 'file_url'}, max_length=limit))
    if 'term' in payload:
        module.term = validate_term(payload.get('term'))
    if 'assessmentType' in payload:
        module.assessment_type = validate_assessment_type(payload.get('assessmentType'))
    for attr, key in [('week', 'week'), ('sequence_order', 'sequenceOrder')]:
        if key in payload:
            setattr(module, attr, parse_int(payload.get(key), 'Sequence field', minimum=1))
    db.commit()
    return {'ok': True}

@teacher_router.delete('/teacher/modules/{module_id}')
def teacher_delete_module(module_id: str, current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    module = db.get(Module, module_id)
    if not module:
        raise HTTPException(status_code=404, detail='Module not found')
    ensure_teacher_course(db, current_user, module.course_id)
    db.delete(module)
    db.commit()
    return {'ok': True}

@teacher_router.post('/teacher/courses/{course_id}/assignments/upload')
async def teacher_upload_assignment_file(course_id: str, file: UploadFile=File(...), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    ensure_teacher_course(db, current_user, course_id)
    upload_dir = os.path.join(os.getcwd(), 'data', 'uploads')
    os.makedirs(upload_dir, exist_ok=True)
    import uuid
    file_id = str(uuid.uuid4())
    ext = os.path.splitext(file.filename)[1]
    safe_filename = f"{file_id}{ext}"
    filepath = os.path.join(upload_dir, safe_filename)
    with open(filepath, "wb") as f:
        f.write(await file.read())
    return {'ok': True, 'filename': file.filename, 'url': f"/api/uploads/{safe_filename}"}

@teacher_router.post('/teacher/courses/{course_id}/assignments')
def teacher_create_assignment(course_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    assignment = Assignment(
        course_id=course.id, 
        title=clean_text(payload.get('title'), max_length=255), 
        instructions=clean_text(payload.get('instructions') or '', required=False, max_length=4000), 
        points=parse_float(payload.get('points') or 100, 'Points', minimum=1), 
        assignment_type=clean_text(payload.get('assignmentType') or 'written_response', max_length=50),
        attachments_json=json.dumps(payload.get('attachments') or []),
        created_by=current_user.id
    )
    db.add(assignment)
    db.flush()
    
    for q in (payload.get('questions') or []):
        db.add(AssignmentQuestion(
            assignment_id=assignment.id,
            section_name=clean_text(q.get('sectionName'), required=False, max_length=120),
            question_text=clean_text(q.get('questionText') or q.get('question'), max_length=2000),
            question_type=clean_text(q.get('type') or 'multiple_choice', max_length=50),
            choices_json=json.dumps(q.get('choices') or []),
            answer_key=json.dumps(q.get('answerKey')) if isinstance(q.get('answerKey'), (list, dict)) else clean_text(q.get('answerKey'), required=False, max_length=1000),
            points=parse_float(q.get('points') or 1, 'Question points', minimum=1)
        ))
        
    db.add(StreamPost(course_id=course.id, author_id=current_user.id, title=assignment.title, body=assignment.instructions, post_type='assignment'))
    db.commit()
    return {'ok': True, 'id': assignment.id}

@teacher_router.put('/teacher/assignments/{assignment_id}')
def teacher_update_assignment(assignment_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if not assignment:
        raise HTTPException(status_code=404, detail='Assignment not found')
    ensure_teacher_course(db, current_user, assignment.course_id)
    for attr, key, limit in [('title', 'title', 255), ('instructions', 'instructions', 4000)]:
        if key in payload:
            setattr(assignment, attr, clean_text(payload.get(key), required=False, max_length=limit))
    if 'points' in payload:
        assignment.points = parse_float(payload.get('points'), 'Points', minimum=1)
    if 'assignmentType' in payload:
        assignment.assignment_type = clean_text(payload.get('assignmentType'), max_length=50)
    if 'attachments' in payload:
        assignment.attachments_json = json.dumps(payload.get('attachments') or [])
    if 'isActive' in payload:
        assignment.is_active = parse_bool(payload.get('isActive'), 'Active status')
        
    if 'questions' in payload:
        from sqlalchemy import delete
        db.execute(delete(AssignmentQuestion).where(AssignmentQuestion.assignment_id == assignment.id))
        for q in payload.get('questions'):
            db.add(AssignmentQuestion(
                assignment_id=assignment.id,
                section_name=clean_text(q.get('sectionName'), required=False, max_length=120),
                question_text=clean_text(q.get('questionText') or q.get('question'), max_length=2000),
                question_type=clean_text(q.get('type') or 'multiple_choice', max_length=50),
                choices_json=json.dumps(q.get('choices') or []),
                answer_key=json.dumps(q.get('answerKey')) if isinstance(q.get('answerKey'), (list, dict)) else clean_text(q.get('answerKey'), required=False, max_length=1000),
                points=parse_float(q.get('points') or 1, 'Question points', minimum=1)
            ))
            
    db.commit()
    return {'ok': True}

@teacher_router.delete('/teacher/assignments/{assignment_id}')
def teacher_delete_assignment(assignment_id: str, current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    assignment = db.get(Assignment, assignment_id)
    if not assignment:
        raise HTTPException(status_code=404, detail='Assignment not found')
    ensure_teacher_course(db, current_user, assignment.course_id)
    assignment.is_active = False
    db.commit()
    return {'ok': True}

@teacher_router.get('/teacher/assignments/{assignment_id}/submissions')
def teacher_assignment_submissions(assignment_id: str, current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> list[dict]:
    assignment = db.get(Assignment, assignment_id)
    if not assignment:
        raise HTTPException(status_code=404, detail='Assignment not found')
    ensure_teacher_course(db, current_user, assignment.course_id)
    
    submissions = db.execute(select(Submission, User).join(User, Submission.student_id == User.id).where(Submission.assignment_id == assignment_id).order_by(User.full_name)).all()
    
    return [{
        'id': sub.id,
        'studentId': user.id,
        'studentName': user.full_name,
        'responseText': sub.response_text,
        'answersJson': sub.answers_json,
        'attachmentsJson': sub.attachments_json,
        'status': sub.status,
        'score': sub.score,
        'feedback': sub.feedback,
        'submittedAt': sub.submitted_at.isoformat() if sub.submitted_at else None
    } for sub, user in submissions]

@teacher_router.post('/teacher/submissions/{submission_id}/grade')
def teacher_grade_submission(submission_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    submission = db.get(Submission, submission_id)
    if not submission:
        raise HTTPException(status_code=404, detail='Submission not found')
        
    assignment = submission.assignment
    course = ensure_teacher_course(db, current_user, assignment.course_id)
    
    score = parse_float(payload.get('score'), 'Score', minimum=0)
    if score > assignment.points:
        raise HTTPException(status_code=400, detail='Score cannot exceed assignment points')
        
    feedback = clean_text(payload.get('feedback'), required=False, max_length=2000)
    
    submission.score = score
    submission.feedback = feedback
    submission.status = 'graded'
    
    grade_entry = db.scalar(select(GradeEntry).where(
        GradeEntry.student_id == submission.student_id,
        GradeEntry.course_id == course.id,
        GradeEntry.component == f"Assignment: {assignment.title[:65]}"
    ))
    
    if grade_entry:
        grade_entry.score = score
        grade_entry.remarks = feedback
    else:
        grade_entry = GradeEntry(
            student_id=submission.student_id,
            course_id=course.id,
            term=course.term,
            component=f"Assignment: {assignment.title[:65]}",
            score=score,
            max_score=assignment.points,
            weight=1.0,
            remarks=feedback,
            recorded_by=current_user.id
        )
        db.add(grade_entry)
        
    db.commit()
    return {'ok': True, 'id': submission.id, 'score': score, 'status': 'graded'}

@teacher_router.post('/teacher/courses/{course_id}/parse-document')
async def teacher_parse_document(course_id: str, action: str=Form('summary'), file: UploadFile=File(...), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    content = await file.read()
    filename = file.filename.lower()
    text = ''
    try:
        if filename.endswith('.pdf'):
            import pypdf
            pdf = pypdf.PdfReader(io.BytesIO(content))
            text = '\\n'.join((page.extract_text() for page in pdf.pages if page.extract_text()))
        elif filename.endswith('.docx'):
            import docx
            doc = docx.Document(io.BytesIO(content))
            text = '\\n'.join((para.text for para in doc.paragraphs))
        elif filename.endswith('.txt') or filename.endswith('.md') or filename.endswith('.csv'):
            text = content.decode('utf-8')
        else:
            raise HTTPException(status_code=400, detail='Unsupported file format. Please upload PDF, DOCX, or TXT.')
    except Exception as e:
        ai_logger.error('Document parsing failed: %s', str(e))
        raise HTTPException(status_code=400, detail='Failed to parse document. Ensure it is not corrupted or password-protected.')
    text = text.strip()
    if not text:
        raise HTTPException(status_code=400, detail='The uploaded document contains no readable text.')
    max_chars = int(os.getenv('DANILO_AI_MAX_DOC_CHARS', '20000'))
    text = text[:max_chars]
    if action == 'summary':
        prompt = f'Summarize the following educational document. Provide the main concepts and key takeaways suitable for students.\\n\\nDocument text:\\n{text}'
    elif action == 'lesson_plan':
        prompt = f'Create a detailed lesson plan based on the following document text. Include Objectives, Subject Matter, Procedure, and Evaluation.\\n\\nDocument text:\\n{text}'
    else:
        prompt = f'Generate 5 key review points and 3 practice questions based on this document.\\n\\nDocument text:\\n{text}'
    try:
        result, metrics = await ask_ollama(SYSTEM_PROMPT_TEACHER, prompt, mode='tutor', memory=None)
    except Exception as e:
        ai_logger.exception("AI generation failed")
        raise HTTPException(status_code=500, detail="AI generation failed.")
    return {'ok': True, 'action': action, 'filename': file.filename, 'result': result, 'metrics': metrics}

@teacher_router.post('/teacher/courses/{course_id}/quizzes/generate')
async def teacher_generate_quiz(course_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    topic = clean_text(payload.get('topic'))
    if not topic:
        raise HTTPException(status_code=400, detail='Topic is required')
    prompt = f'You are an expert DepEd teacher. Generate a 5-item multiple-choice quiz based on the following topic.\nEach question must have exactly 4 options.\nOutput the quiz in strict JSON format.\n\nTopic: {topic}\nSubject: {course.subject}\nGrade Level: {course.grade_level}\n\nJSON Schema:\n{{\n  "title": "Generated Quiz",\n  "questions": [\n    {{\n      "question": "Question text here?",\n      "choices": ["A) Option 1", "B) Option 2", "C) Option 3", "D) Option 4"],\n      "answerKey": "A) Option 1"\n    }}\n  ]\n}}\n\nReturn ONLY valid JSON. Do not include markdown formatting or explanations.\n'
    result, metrics = await ask_ollama(prompt, mode='tutor', memory=None)
    try:
        import re
        json_str = result
        if '```json' in json_str:
            json_str = json_str.split('```json')[1].split('```')[0]
        elif '```' in json_str:
            json_str = json_str.split('```')[1].split('```')[0]
        json_str = json_str.strip()
        try:
            data = json.loads(json_str)
        except json.JSONDecodeError:
            match = re.search(r'\{.*\}', json_str, re.DOTALL)
            if match:
                data = json.loads(match.group(0))
            else:
                raise
        return {'ok': True, 'quiz': data, 'metrics': metrics}
    except Exception as e:
        ai_logger.error('Failed to parse quiz JSON: %s. Response was: %s', str(e), result)
        raise HTTPException(status_code=500, detail='AI failed to generate a valid quiz format. Please try again.')

@teacher_router.post('/teacher/courses/{course_id}/quizzes/generate-from-file')
async def teacher_generate_quiz_from_file(
    course_id: str, 
    file: UploadFile=File(...), 
    questionCount: str=Form('10'),
    difficulty: str=Form('medium'),
    questionTypes: str=Form('multiple_choice'),
    current_user: User=Depends(get_current_user), 
    db: Session=Depends(get_db)
) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    content = await file.read()
    filename = file.filename.lower()
    text = ''
    try:
        if filename.endswith('.pdf'):
            import pypdf
            pdf = pypdf.PdfReader(io.BytesIO(content))
            text = '\\n'.join((page.extract_text() for page in pdf.pages if page.extract_text()))
        elif filename.endswith('.docx'):
            import docx
            doc = docx.Document(io.BytesIO(content))
            text = '\\n'.join((para.text for para in doc.paragraphs))
        elif filename.endswith('.txt') or filename.endswith('.md') or filename.endswith('.csv'):
            text = content.decode('utf-8')
        else:
            raise HTTPException(status_code=400, detail='Unsupported file format. Please upload PDF, DOCX, or TXT.')
    except Exception as e:
        ai_logger.error('Document parsing failed: %s', str(e))
        raise HTTPException(status_code=400, detail='Failed to parse document.')
    
    text = text.strip()[:20000]
    if not text:
        raise HTTPException(status_code=400, detail='No readable text found.')

    prompt = f'''You are an expert DepEd teacher. Generate a comprehensive {questionCount}-item quiz based on the following document.
Difficulty: {difficulty}
Include the following question types: {questionTypes}.
Output the quiz in strict JSON format.

Document Text: {text}
Subject: {course.subject}
Grade Level: {course.grade_level}

JSON Schema:
{{
  "title": "Generated Quiz from Document",
  "questions": [
    {{
      "type": "multiple_choice", // Must be one of: multiple_choice, checkbox, short_answer, identification, true_false
      "question": "Question text here?",
      "choices": ["Option 1", "Option 2"], // required for multiple_choice, checkbox, true_false
      "answerKey": "Correct answer text", // string for multiple_choice/identification/short_answer/true_false, array of strings for checkbox
      "points": 1
    }}
  ]
}}

Return ONLY valid JSON. Do not include markdown formatting or explanations.
'''
    result, metrics = await ask_ollama(prompt, mode='tutor', memory=None)
    try:
        import re
        json_str = result
        if '```json' in json_str:
            json_str = json_str.split('```json')[1].split('```')[0]
        elif '```' in json_str:
            json_str = json_str.split('```')[1].split('```')[0]
        json_str = json_str.strip()
        try:
            data = json.loads(json_str)
        except json.JSONDecodeError:
            match = re.search(r'\{.*\}', json_str, re.DOTALL)
            if match:
                data = json.loads(match.group(0))
            else:
                raise
        return {'ok': True, 'quiz': data, 'metrics': metrics}
    except Exception as e:
        ai_logger.error('Failed to parse quiz JSON: %s. Response was: %s', str(e), result)
        raise HTTPException(status_code=500, detail='AI failed to generate a valid quiz format. Please try again.')

@teacher_router.post('/teacher/courses/{course_id}/quizzes')
def teacher_create_quiz(course_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    quiz = Quiz(course_id=course.id, title=clean_text(payload.get('title') or 'Class Quiz', max_length=255), instructions=clean_text(payload.get('instructions') or 'Answer each question based on the current lesson.', max_length=2000), is_published=parse_bool(payload.get('isPublished'), 'Published status', default=False), created_by=current_user.id)
    db.add(quiz)
    db.flush()
    for item in payload.get('questions') or []:
        db.add(QuizQuestion(quiz_id=quiz.id, question_text=clean_text(item.get('questionText') or item.get('question'), max_length=2000), choices_json=json.dumps(item.get('choices') or []) if 'choices' in item else clean_text(item.get('choicesJson'), required=False, max_length=4000), answer_key=json.dumps(item.get('answerKey')) if isinstance(item.get('answerKey'), (list, dict)) else clean_text(item.get('answerKey'), required=False, max_length=1000), points=parse_float(item.get('points') or 1, 'Question points', minimum=1)))
    db.commit()
    return {'ok': True, 'id': quiz.id}

@teacher_router.get('/teacher/courses/{course_id}/gradebook')
def teacher_gradebook(course_id: str, current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    course = ensure_teacher_course(db, current_user, course_id)
    students = teacher_course_students(course_id, current_user, db)
    grades = db.execute(select(GradeEntry, User).join(User, GradeEntry.student_id == User.id).where(GradeEntry.course_id == course_id).order_by(User.full_name, GradeEntry.created_at)).all()
    by_student: dict[int, list[GradeEntry]] = {}
    names: dict[int, str] = {}
    for grade, student in grades:
        by_student.setdefault(student.id, []).append(grade)
        names[student.id] = student.full_name
    summaries = []
    for student_id, rows in by_student.items():
        for item in summarize_grade_entries(course, rows, names.get(student_id)):
            item['studentId'] = student_id
            summaries.append(item)
    return {'course': serialize_course(course), 'students': students, 'entries': [{'id': grade.id, 'studentId': student.id, 'studentName': student.full_name, 'term': grade.term, 'component': grade.component, 'score': grade.score, 'maxScore': grade.max_score, 'weight': grade.weight, 'remarks': grade.remarks or ''} for grade, student in grades], 'grades': summaries}

@teacher_router.get('/teacher/insights')
async def teacher_student_insights(class_id: str | None=None, subject: str | None=None, include_ai: bool=True, current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    stmt = select(Course).where(Course.is_active == True, Course.teacher_id == current_user.id)
    if subject:
        stmt = stmt.where(Course.subject == subject)
    if class_id:
        stmt = stmt.where(Course.id == class_id)
    course = db.scalars(stmt.order_by(Course.subject.asc(), Course.title.asc())).first()
    if not course:
        return {'course': None, 'students': [], 'strugglingStudents': [], 'classWeakTopics': [], 'stats': {'studentCount': 0, 'strugglingCount': 0, 'atRiskCount': 0, 'needsAttentionCount': 0, 'classAverage': None, 'missingSubmissions': 0}, 'aiSummary': 'No class data available for the selected filters.', 'aiStatus': 'skipped'}
    analysis = analyze_student_performance(db, course.id)
    analysis['availableClasses'] = build_admin_course_cards(db) if current_user.role == 'admin' else build_teacher_course_cards(db, current_user.id)
    analysis['availableSubjects'] = sorted({item['subject'] for item in analysis['availableClasses'] if item.get('subject')})
    analysis['aiSummary'] = ''
    analysis['aiStatus'] = 'skipped'
    if include_ai:
        try:
            ai_summary, metrics = await ask_ollama(build_student_insights_prompt(analysis), 'short')
            analysis['aiSummary'] = ai_summary or 'DANILO did not return additional recommendations.'
            analysis['aiStatus'] = 'ready'
            analysis['aiMetrics'] = metrics
        except Exception:
            logger.exception('Ollama failed while generating student insights for course_id=%s', course.id)
            analysis['aiSummary'] = 'DANILO Tutor is offline or still getting ready. Deterministic insights are shown below.'
            analysis['aiStatus'] = 'offline'
    return analysis

@teacher_router.post('/teacher/courses/{course_id}/grades')
def teacher_create_grade(course_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    ensure_teacher_course(db, current_user, course_id)
    student_id = parse_id(payload.get('studentId') or payload.get('student_id'), 'Student')
    if not db.scalar(select(Enrollment).where(Enrollment.course_id == course_id, Enrollment.student_id == student_id, Enrollment.status == 'active')):
        raise HTTPException(status_code=400, detail='Student is not enrolled in this class')
    grade = GradeEntry(student_id=student_id, course_id=course_id, term=validate_term(payload.get('term') or 'Term 1'), component=clean_text(payload.get('component'), max_length=80), score=parse_float(payload.get('score'), 'Score', minimum=0), max_score=parse_float(payload.get('maxScore') or payload.get('max_score') or 100, 'Max score', minimum=1), weight=parse_float(payload.get('weight') or 1, 'Weight', minimum=0), remarks=clean_text(payload.get('remarks'), required=False, max_length=1000), recorded_by=current_user.id)
    db.add(grade)
    db.commit()
    return {'ok': True, 'id': grade.id}

@teacher_router.put('/teacher/grades/{grade_id}')
def teacher_update_grade(grade_id: str, payload: dict=Body(default={}), current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    grade = db.get(GradeEntry, grade_id)
    if not grade:
        raise HTTPException(status_code=404, detail='Grade not found')
    ensure_teacher_course(db, current_user, grade.course_id)
    for attr, key in [('term', 'term'), ('component', 'component'), ('remarks', 'remarks')]:
        if key in payload:
            if attr == 'term':
                grade.term = validate_term(payload.get(key))
            else:
                setattr(grade, attr, clean_text(payload.get(key), required=attr != 'remarks', max_length=1000 if attr == 'remarks' else 80))
    for attr, key in [('score', 'score'), ('max_score', 'maxScore'), ('weight', 'weight')]:
        if key in payload:
            minimum = 1 if attr == 'max_score' else 0
            setattr(grade, attr, parse_float(payload.get(key), key, minimum=minimum))
    db.commit()
    return {'ok': True}

@teacher_router.delete('/teacher/grades/{grade_id}')
def teacher_delete_grade(grade_id: str, current_user: User=Depends(get_current_user), db: Session=Depends(get_db)) -> dict:
    grade = db.get(GradeEntry, grade_id)
    if not grade:
        raise HTTPException(status_code=404, detail='Grade not found')
    ensure_teacher_course(db, current_user, grade.course_id)
    db.delete(grade)
    db.commit()
    return {'ok': True}

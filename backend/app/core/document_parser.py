import os
import re

def extract_text_from_file(file_path: str, mime_type: str) -> str:
    text = ""
    try:
        if "pdf" in mime_type or file_path.endswith('.pdf'):
            import pypdf
            with open(file_path, "rb") as f:
                reader = pypdf.PdfReader(f)
                text = "\\n".join([page.extract_text() for page in reader.pages if page.extract_text()])
                
        elif "wordprocessingml" in mime_type or file_path.endswith('.docx'):
            import docx
            doc = docx.Document(file_path)
            text = "\\n".join([para.text for para in doc.paragraphs])
            
        elif "presentationml" in mime_type or file_path.endswith('.pptx'):
            import pptx
            prs = pptx.Presentation(file_path)
            slides_text = []
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        slides_text.append(shape.text)
            text = "\\n".join(slides_text)
            
        elif "text/plain" in mime_type or file_path.endswith('.txt'):
            with open(file_path, "r", encoding="utf-8") as f:
                text = f.read()
    except Exception as e:
        print(f"Failed to parse {file_path}: {e}")
        
    return text.strip()

def chunk_text(text: str, words_per_chunk: int = 300, overlap_words: int = 50) -> list[str]:
    words = text.split()
    chunks = []
    
    if not words:
        return chunks
        
    start = 0
    while start < len(words):
        end = start + words_per_chunk
        chunk = " ".join(words[start:end])
        chunks.append(chunk)
        if end >= len(words):
            break
        start += (words_per_chunk - overlap_words)
        
    return chunks
